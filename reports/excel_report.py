"""
Excel Report Generator Module.
Creates multi-sheet financial Excel workbooks with corporate styling, Data Quality metrics, conditional formatting, and embedded charts using OpenPyXL.
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analytics.kpi_analysis import KPIAnalytics
from analytics.market_analysis import MarketAnalytics
from config.config import Config, get_config
from database.db_connection import DatabaseManager

logger = logging.getLogger("MarketPipeline.ExcelReport")


class ExcelReportGenerator:
    """Generates corporate-grade Excel reports with styling and charts."""

    def __init__(self, config: Optional[Config] = None, db: Optional[DatabaseManager] = None):
        self.config = config or get_config()
        self.db = db
        self.market_analytics = MarketAnalytics(self.config, self.db)
        self.kpi_analytics = KPIAnalytics(self.config, self.db)

        # Style Definitions
        self.header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        self.header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        
        self.sub_header_fill = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
        self.sub_header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

        self.card_header_fill = PatternFill(start_color="E6F0FA", end_color="E6F0FA", fill_type="solid")
        self.card_header_font = Font(name="Segoe UI", size=10, bold=True, color="1B365D")

        self.zebra_fill = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
        self.green_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
        self.green_font = Font(name="Segoe UI", size=10, color="276A3C", bold=True)
        self.red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        self.red_font = Font(name="Segoe UI", size=10, color="C00000", bold=True)

        self.regular_font = Font(name="Segoe UI", size=10)
        self.bold_font = Font(name="Segoe UI", size=10, bold=True)

        thin_border_side = Side(border_style="thin", color="D3D3D3")
        self.cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    def generate_full_report(
        self,
        output_filename: Optional[str] = None,
        data_quality_score: Optional[float] = None,
    ) -> Path:
        """Generate full 3-sheet market workbook."""
        logger.info("================== [PHASE 5: EXCEL REPORT GENERATION] ==================")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        df_market = self.market_analytics.get_latest_market_data()
        df_sector = self.market_analytics.calculate_sector_performance(df_market)
        kpi_summary = self.kpi_analytics.get_executive_summary_kpis(df_market)
        breadth = self.market_analytics.calculate_market_breadth(df_market)

        # 1. Sheet 1: Market Data
        ws_market = wb.create_sheet(title="Market Data")
        self._build_market_data_sheet(ws_market, df_market)

        # 2. Sheet 2: KPI Summary
        ws_kpi = wb.create_sheet(title="KPI Summary")
        self._build_kpi_summary_sheet(ws_kpi, df_market, kpi_summary, breadth, data_quality_score)

        # 3. Sheet 3: Sector Analysis
        ws_sector = wb.create_sheet(title="Sector Analysis")
        self._build_sector_analysis_sheet(ws_sector, df_sector)

        # Auto-adjust column widths
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val = str(cell.value or "")
                    if cell.number_format and "%" in cell.number_format:
                        val = f"{val}%"
                    max_len = max(max_len, len(val))
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Save workbook
        if not output_filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"market_analytics_report_{timestamp}.xlsx"

        output_path = self.config.REPORTS_DIR / output_filename
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        logger.info("Excel financial report successfully generated at: %s", output_path)
        return output_path

    def _build_market_data_sheet(self, ws, df):
        """Construct Sheet 1: Full market data table with formatting."""
        headers = [
            "Symbol", "Company Name", "Sector", "Industry", "Close Price (₹)",
            "Change %", "Volume", "Market Cap (Cr ₹)", "P/E Ratio"
        ]
        
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, r in enumerate(df.itertuples(), start=2):
            row_data = [
                r.symbol,
                r.company_name,
                r.sector,
                r.industry,
                r.close_price,
                (r.change_percent / 100.0) if r.change_percent is not None else None,
                r.volume,
                getattr(r, "market_cap", None),
                getattr(r, "pe_ratio", None),
            ]
            ws.append(row_data)

            is_even = (row_idx % 2 == 0)
            fill = self.zebra_fill if is_even else PatternFill(fill_type=None)

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = self.regular_font
                cell.border = self.cell_border
                cell.fill = fill

                if col_idx == 5:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx == 6:
                    cell.number_format = '+0.00%;-0.00%;0.00%'
                    cell.alignment = Alignment(horizontal="right")
                    if r.change_percent is not None:
                        if r.change_percent > 0:
                            cell.fill = self.green_fill
                            cell.font = self.green_font
                        elif r.change_percent < 0:
                            cell.fill = self.red_fill
                            cell.font = self.red_font
                elif col_idx in [7, 8]:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx == 9:
                    cell.number_format = '0.00'
                    cell.alignment = Alignment(horizontal="right")

    def _build_kpi_summary_sheet(self, ws, df, kpis, breadth, data_quality_score=None):
        """Construct Sheet 2: Executive KPI Cards & Top Gainers/Losers."""
        ws.merge_cells("A1:E1")
        title_cell = ws.cell(row=1, column=1, value="EXECUTIVE MARKET KPI SUMMARY")
        title_cell.fill = self.header_fill
        title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=2, column=1, value=f"Last Updated: {datetime.now().strftime('%d %B %Y, %H:%M:%S')}")
        ws.cell(row=2, column=1).font = Font(name="Segoe UI", size=9, italic=True)

        dq_display = f"{data_quality_score:.1f}%" if data_quality_score is not None else "100.0%"

        kpi_rows = [
            ("Total Companies Tracked", kpis["total_companies"], "Count"),
            ("Advancing Stocks", breadth["advancers"], "Gainers Count"),
            ("Declining Stocks", breadth["decliners"], "Losers Count"),
            ("Market Breadth (A/D Ratio)", f"{breadth['ad_ratio']} ({breadth['market_sentiment']})", f"{breadth['advancers']} Up / {breadth['decliners']} Down"),
            ("Average Daily Return", f"{kpis['avg_daily_return']:+.2f}%", "Market Average"),
            ("Best Performer", f"{kpis['top_gainer_symbol']} ({kpis['top_gainer_change']:+.2f}%)", "Top Daily Gainer"),
            ("Worst Performer", f"{kpis['top_loser_symbol']} ({kpis['top_loser_change']:+.2f}%)", "Top Daily Loser"),
            ("Data Quality Score", dq_display, "Weighted Pipeline Integrity Score"),
            ("Total Market Cap (Cr)", f"₹{kpis['total_market_cap_cr']:,.2f} Cr", "Aggregate Market Cap"),
            ("Average P/E Ratio", f"{kpis['avg_pe_ratio']:.2f}", "Valuation Multiple"),
        ]

        ws.cell(row=4, column=1, value="Metric").fill = self.sub_header_fill
        ws.cell(row=4, column=1).font = self.sub_header_font
        ws.cell(row=4, column=2, value="Value").fill = self.sub_header_fill
        ws.cell(row=4, column=2).font = self.sub_header_font
        ws.cell(row=4, column=3, value="Notes").fill = self.sub_header_fill
        ws.cell(row=4, column=3).font = self.sub_header_font

        for idx, (metric, val, note) in enumerate(kpi_rows, start=5):
            c1 = ws.cell(row=idx, column=1, value=metric)
            c2 = ws.cell(row=idx, column=2, value=val)
            c3 = ws.cell(row=idx, column=3, value=note)
            for c in (c1, c2, c3):
                c.font = self.regular_font
                c.border = self.cell_border
            c1.font = self.bold_font

        # Top 5 Gainers Table
        start_row = 17
        ws.cell(row=start_row, column=1, value="TOP 5 GAINERS").font = Font(name="Segoe UI", size=11, bold=True, color="276A3C")
        gainers_headers = ["Symbol", "Company", "Sector", "Price (₹)", "Change %"]
        for c_idx, h in enumerate(gainers_headers, start=1):
            cell = ws.cell(row=start_row + 1, column=c_idx, value=h)
            cell.fill = self.card_header_fill
            cell.font = self.card_header_font

        top_gainers = self.kpi_analytics.get_top_gainers(5, df)
        for g_idx, r in enumerate(top_gainers.itertuples(), start=start_row + 2):
            ws.cell(row=g_idx, column=1, value=r.symbol).border = self.cell_border
            ws.cell(row=g_idx, column=2, value=r.company_name).border = self.cell_border
            ws.cell(row=g_idx, column=3, value=r.sector).border = self.cell_border
            p_cell = ws.cell(row=g_idx, column=4, value=r.close_price)
            p_cell.number_format = '₹#,##0.00'
            p_cell.border = self.cell_border
            chg_cell = ws.cell(row=g_idx, column=5, value=r.change_percent / 100.0)
            chg_cell.number_format = '+0.00%'
            chg_cell.fill = self.green_fill
            chg_cell.font = self.green_font
            chg_cell.border = self.cell_border

        # Top 5 Losers Table
        losers_start = start_row + 9
        ws.cell(row=losers_start, column=1, value="TOP 5 LOSERS").font = Font(name="Segoe UI", size=11, bold=True, color="C00000")
        for c_idx, h in enumerate(gainers_headers, start=1):
            cell = ws.cell(row=losers_start + 1, column=c_idx, value=h)
            cell.fill = self.card_header_fill
            cell.font = self.card_header_font

        top_losers = self.kpi_analytics.get_top_losers(5, df)
        for l_idx, r in enumerate(top_losers.itertuples(), start=losers_start + 2):
            ws.cell(row=l_idx, column=1, value=r.symbol).border = self.cell_border
            ws.cell(row=l_idx, column=2, value=r.company_name).border = self.cell_border
            ws.cell(row=l_idx, column=3, value=r.sector).border = self.cell_border
            p_cell = ws.cell(row=l_idx, column=4, value=r.close_price)
            p_cell.number_format = '₹#,##0.00'
            p_cell.border = self.cell_border
            chg_cell = ws.cell(row=l_idx, column=5, value=r.change_percent / 100.0)
            chg_cell.number_format = '-0.00%'
            chg_cell.fill = self.red_fill
            chg_cell.font = self.red_font
            chg_cell.border = self.cell_border

    def _build_sector_analysis_sheet(self, ws, df_sector):
        """Construct Sheet 3: Sector metrics table and embedded OpenPyXL Chart."""
        headers = ["Sector", "Number of Stocks", "Average Return %", "Sector Volatility %", "Total Market Cap (Cr ₹)", "Total Volume", "Advancers", "Decliners"]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, r in enumerate(df_sector.itertuples(), start=2):
            row_data = [
                r.sector,
                r.company_count,
                r.avg_return_pct / 100.0,
                getattr(r, "sector_volatility", 0.0) / 100.0,
                r.total_market_cap_cr,
                r.total_volume,
                r.advancers,
                r.decliners,
            ]
            ws.append(row_data)

            is_even = (row_idx % 2 == 0)
            fill = self.zebra_fill if is_even else PatternFill(fill_type=None)

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = self.regular_font
                cell.border = self.cell_border
                cell.fill = fill
                if col_idx in [3, 4]:
                    cell.number_format = '+0.00%;-0.00%;0.00%'
                    cell.alignment = Alignment(horizontal="right")
                    if col_idx == 3:
                        if r.avg_return_pct > 0:
                            cell.fill = self.green_fill
                            cell.font = self.green_font
                        elif r.avg_return_pct < 0:
                            cell.fill = self.red_fill
                            cell.font = self.red_font
                elif col_idx == 5:
                    cell.number_format = '₹#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx in [2, 6, 7, 8]:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="right")

        if len(df_sector) > 0:
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "Sector Average Daily Return (%)"
            chart.y_axis.title = "Return (%)"
            chart.x_axis.title = "Sector"
            chart.height = 12
            chart.width = 18

            data = Reference(ws, min_col=3, min_row=1, max_row=len(df_sector) + 1)
            categories = Reference(ws, min_col=1, min_row=2, max_row=len(df_sector) + 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            chart.legend = None

            ws.add_chart(chart, "J2")


def generate_excel_report(
    valid_data: Optional[Any] = None,
    analytics: Optional[Any] = None,
    data_quality_score: Optional[float] = None,
    output_filename: Optional[str] = None,
    config: Optional[Config] = None,
) -> Path:
    """Functional helper for Excel report generation."""
    generator = ExcelReportGenerator(config)
    return generator.generate_full_report(output_filename, data_quality_score=data_quality_score)
